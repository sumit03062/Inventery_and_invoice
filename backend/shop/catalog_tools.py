import csv
import hashlib
import io
import zipfile
from django.db import transaction
from django.http import HttpResponse
from rest_framework.decorators import api_view
from rest_framework.exceptions import ValidationError
from rest_framework.response import Response
from .models import Product, ProductImport
from .serializers import ProductSerializer
from .permissions import require
from . import services as s

FIELDS=['name','sku','barcode','unit','pieces_per_box','purchase_price','price','gst_percent','hsn_code','opening_stock','low_stock_threshold']


@api_view(['GET'])
def import_template(request):
    require(request.user,'products.write')
    stream=io.StringIO()
    writer=csv.writer(stream)
    writer.writerow(FIELDS)
    writer.writerow(['Example product','SKU-001','','PCS',12,'20','30','0','','48','5'])
    response=HttpResponse(stream.getvalue(),content_type='text/csv')
    response['Content-Disposition']='attachment; filename="products-template.csv"'
    return response


def read_rows(file):
    if not file or file.size>5*1024*1024:
        raise ValidationError('Choose a CSV or XLSX file under 5 MB.')
    content=file.read()
    try:
        if file.name.lower().endswith('.xlsx'):
            from openpyxl import load_workbook
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(x.file_size for x in archive.infolist())>20*1024*1024:
                    raise ValidationError('Expanded workbook is too large.')
            workbook=load_workbook(io.BytesIO(content),read_only=True,data_only=False,keep_links=False)
            try:
                iterator=workbook.active.iter_rows(values_only=True)
                headers=[str(x or '').strip() for x in next(iterator)]
                rows=[]
                for values in iterator:
                    if all(x is None for x in values):
                        continue
                    if any(isinstance(x,str) and x.startswith('=') for x in values):
                        raise ValidationError('Use values, not formulas, in the workbook.')
                    rows.append(dict(zip(headers,values)))
                    if len(rows)>500:
                        break
            finally:
                workbook.close()
        elif file.name.lower().endswith('.csv'):
            reader=csv.DictReader(io.StringIO(content.decode('utf-8-sig')))
            headers=reader.fieldnames or []
            rows=[]
            for row in reader:
                if any(row.values()):
                    rows.append(row)
                if len(rows)>500:
                    break
        else:
            raise ValidationError('Only CSV and XLSX files are supported.')
        if not {'name','sku','price'}.issubset(headers) or not 1<=len(rows)<=500:
            raise ValidationError('Include name, sku and price columns and 1–500 product rows.')
        return rows,hashlib.sha256(content).hexdigest()
    except (UnicodeError,ValueError,KeyError,StopIteration,zipfile.BadZipFile) as exc:
        raise ValidationError('Unable to read the file. Use the import template.') from exc


@api_view(['POST'])
@transaction.atomic
def import_products(request):
    require(request.user,'products.write')
    require(request.user,'prices.write')
    require(request.user,'inventory.adjust')
    s.lock_shop()
    rows,digest=read_rows(request.FILES.get('file'))
    commit=request.data.get('commit')=='true'
    key=s.request_key(request.data) if commit else None
    old=ProductImport.objects.filter(request_key=key).first() if key else None
    if old:
        if old.digest!=digest or old.created_by_id!=request.user.id:
            raise ValidationError('This import key belongs to another file.')
        return Response({'created':old.count})
    prepared=[]
    errors=[]
    skus=set()
    barcodes=set()
    for n,row in enumerate(rows,2):
        data={k:v for k,v in row.items() if k in FIELDS and v is not None and str(v).strip()!=''}
        for k in ['sku','barcode','hsn_code']:
            if k in data:
                data[k]=str(data[k]).strip()
        try:
            serializer=ProductSerializer(data=data,context={'request':request})
            serializer.is_valid(raise_exception=True)
            sku=serializer.validated_data['sku']
            barcode=serializer.validated_data.get('barcode')
            if sku in skus or barcode and barcode in barcodes:
                raise ValidationError('Duplicate SKU or barcode within file.')
            skus.add(sku)
            if barcode:
                barcodes.add(barcode)
            unit=serializer.validated_data.get('unit','PCS')
            opening=s.quantity(data.get('opening_stock',0),unit,minimum=0)
            prepared.append((serializer,opening))
        except ValidationError as exc:
            errors.append({'row':n,'error':exc.detail})
    if errors:
        return Response({'errors':errors,'created':0},status=400)
    if not commit:
        return Response({'valid':True,'count':len(prepared),'preview':[{'name':x.validated_data['name'],'sku':x.validated_data['sku'],'opening_stock':qty} for x,qty in prepared]})
    for serializer,opening in prepared:
        p=serializer.save()
        if opening:
            s.stock(p,opening,request.user,'Imported opening stock')
    ProductImport.objects.create(request_key=key,digest=digest,count=len(prepared),created_by=request.user)
    s.audit(request.user,'PRODUCTS_IMPORTED',f'{len(prepared)} products')
    return Response({'created':len(prepared)},status=201)


@api_view(['GET'])
def labels(request):
    require(request.user,'products.write')
    from reportlab.graphics.barcode import createBarcodeDrawing
    from reportlab.platypus import Spacer,Table
    from .documents import document,p
    ids=str(request.query_params.get('ids','')).split(',')
    if not 1<=len(ids)<=100 or any(not x.isdigit() for x in ids):
        raise ValidationError('Select 1–100 products.')
    copies=s.integer(request.query_params.get('copies',1),maximum=20)
    products=list(Product.objects.filter(id__in=ids,active=True))
    if len(products)!=len(set(ids)) or len(products)*copies>200:
        raise ValidationError('Choose active products and at most 200 labels.')
    cells=[]
    for product in products:
        code=product.barcode or product.sku
        if not code.isascii() or not code.isprintable() or len(code)>48:
            raise ValidationError(f'{product.name}: use a printable ASCII barcode or SKU of up to 48 characters.')
        for _ in range(copies):
            drawing=createBarcodeDrawing('Code128',value=code,barHeight=32,humanReadable=True)
            scale=min(1,210/drawing.width)
            drawing.scale(scale,scale)
            drawing.width*=scale
            drawing.height*=scale
            cells.append([p(product.name),p(f'INR {product.price} / {product.unit}'),drawing,Spacer(1,12)])
    if len(cells)%2:
        cells.append('')
    rows=[cells[i:i+2] for i in range(0,len(cells),2)]
    return document([p('Product labels','Title'),Table(rows,colWidths=[240,240])],'product-labels')
